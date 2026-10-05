from typing import List, TypedDict
from pathlib import Path
from urllib.parse import urlparse
import os
import re
import plistlib
import tempfile

import requests
from bs4 import BeautifulSoup
from pydantic import BaseModel, Field
from dotenv import load_dotenv

from langchain_community.document_loaders import PyPDFLoader
from langchain_community.vectorstores import FAISS
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_groq import ChatGroq
from langchain_tavily import TavilySearch
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_core.prompts import ChatPromptTemplate

from langgraph.graph import StateGraph, START, END


# ============================================================
# ENVIRONMENT
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
DOCUMENTS_DIR = BASE_DIR / "documents"
ENV_FILE = BASE_DIR / ".env"

load_dotenv(dotenv_path=ENV_FILE)

if not os.getenv("GROQ_API_KEY"):
    raise RuntimeError(
        f"GROQ_API_KEY was not found.\n"
        f"Create this file:\n{ENV_FILE}\n\n"
        f"and add:\n"
        f"GROQ_API_KEY=your_groq_api_key"
    )

if not os.getenv("TAVILY_API_KEY"):
    raise RuntimeError(
        f"TAVILY_API_KEY was not found.\n"
        f"Create this file:\n{ENV_FILE}\n\n"
        f"and add:\n"
        f"TAVILY_API_KEY=your_tavily_api_key"
    )


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text: str) -> str:
    text = text.encode("utf-8", "ignore").decode("utf-8", "ignore")
    text = re.sub(r"\s+", " ", text).strip()
    return text


# ============================================================
# URL HELPERS
# ============================================================

URL_PATTERN = re.compile(r"https?://[^\s<>\"]+")


def normalize_url(url: str) -> str:
    return url.strip().rstrip(".,;:)]}")


def extract_urls(text: str) -> List[str]:
    return list(dict.fromkeys(
        normalize_url(url)
        for url in URL_PATTERN.findall(text)
    ))


def is_valid_url(url: str) -> bool:
    try:
        parsed = urlparse(url)
        return parsed.scheme in {"http", "https"} and bool(parsed.netloc)
    except Exception:
        return False


# ============================================================
# LOCAL FILE LOADERS
# ============================================================

def load_pdf_file(path: Path) -> List[Document]:
    try:
        docs = PyPDFLoader(str(path)).load()
    except Exception as e:
        print(f"[WARNING] Could not read PDF {path}: {e}")
        return []

    output = []

    for doc in docs:
        text = clean_text(doc.page_content)
        if not text:
            continue

        doc.page_content = text
        doc.metadata.update(
            {
                "source": str(path),
                "type": "file",
                "file_type": "pdf",
            }
        )
        output.append(doc)

    return output


def load_docx_file(path: Path) -> List[Document]:
    try:
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        paragraphs = [
            paragraph.text.strip()
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ]
        text = clean_text("\n".join(paragraphs))

    except Exception as e:
        print(f"[WARNING] Could not read DOCX {path}: {e}")
        return []

    if not text:
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": str(path),
                "type": "file",
                "file_type": "docx",
            },
        )
    ]


def load_text_file(path: Path) -> List[Document]:
    try:
        text = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )
    except Exception as e:
        print(f"[WARNING] Could not read {path}: {e}")
        return []

    text = clean_text(text)

    if not text:
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": str(path),
                "type": "file",
                "file_type": path.suffix.lower(),
            },
        )
    ]


# ============================================================
# LINK FILE LOADERS
# ============================================================

def extract_url_from_shortcut(path: Path) -> str | None:
    try:
        suffix = path.suffix.lower()

        if suffix == ".url":
            text = path.read_text(
                encoding="utf-8",
                errors="ignore",
            )
            match = re.search(
                r"(?im)^\s*URL\s*=\s*(https?://\S+)\s*$",
                text,
            )
            return (
                normalize_url(match.group(1))
                if match
                else None
            )

        if suffix == ".webloc":
            data = plistlib.loads(path.read_bytes())
            url = data.get("URL")
            return (
                normalize_url(url)
                if isinstance(url, str)
                else None
            )

    except Exception as e:
        print(f"[WARNING] Could not read link shortcut {path}: {e}")

    return None


def extract_urls_from_link_file(path: Path) -> List[str]:
    try:
        text = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )
        return [
            url
            for url in extract_urls(text)
            if is_valid_url(url)
        ]
    except Exception as e:
        print(f"[WARNING] Could not read link file {path}: {e}")
        return []


# ============================================================
# FETCH URL
# ============================================================

def fetch_url(url: str) -> List[Document]:
    if not is_valid_url(url):
        return []

    try:
        response = requests.get(
            url,
            timeout=20,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 "
                    "(KHTML, like Gecko) "
                    "Chrome/154 Safari/537.36"
                )
            },
        )
        response.raise_for_status()

        content_type = (
            response.headers.get("Content-Type", "").lower()
        )

        # ----------------------------------------------------
        # PDF URL
        # ----------------------------------------------------

        if (
            "application/pdf" in content_type
            or url.lower().split("?")[0].endswith(".pdf")
        ):
            from pypdf import PdfReader

            temp_path = None

            try:
                with tempfile.NamedTemporaryFile(
                    suffix=".pdf",
                    delete=False,
                ) as temp:
                    temp.write(response.content)
                    temp_path = temp.name

                reader = PdfReader(temp_path)
                docs = []

                for page_number, page in enumerate(
                    reader.pages,
                    start=1,
                ):
                    text = clean_text(
                        page.extract_text() or ""
                    )

                    if not text:
                        continue

                    docs.append(
                        Document(
                            page_content=text,
                            metadata={
                                "source": url,
                                "url": url,
                                "type": "url",
                                "file_type": "pdf",
                                "page": page_number,
                            },
                        )
                    )

                return docs

            finally:
                if temp_path:
                    try:
                        Path(temp_path).unlink(missing_ok=True)
                    except Exception:
                        pass

        # ----------------------------------------------------
        # HTML URL
        # ----------------------------------------------------

        soup = BeautifulSoup(
            response.text,
            "html.parser",
        )

        for tag in soup(
            ["script", "style", "noscript", "svg"]
        ):
            tag.decompose()

        title = (
            soup.title.get_text(" ", strip=True)
            if soup.title
            else url
        )

        text = clean_text(
            soup.get_text(" ", strip=True)
        )

        if not text:
            return []

        return [
            Document(
                page_content=(
                    f"TITLE: {title}\n"
                    f"URL: {url}\n"
                    f"CONTENT:\n{text}"
                ),
                metadata={
                    "source": url,
                    "url": url,
                    "type": "url",
                    "title": title,
                },
            )
        ]

    except Exception as e:
        print(f"[WARNING] Could not fetch {url}: {e}")
        return []


# ============================================================
# LOAD DOCUMENTS DIRECTORY
# ============================================================

def load_documents_from_folder(
    folder: Path,
) -> List[Document]:

    if not folder.exists():
        raise FileNotFoundError(
            f"Documents folder not found:\n{folder}"
        )

    if not folder.is_dir():
        raise NotADirectoryError(
            f"Not a directory:\n{folder}"
        )

    supported_files = {
        ".pdf",
        ".docx",
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".html",
        ".htm",
        ".url",
        ".webloc",
        ".links",
        ".urls",
    }

    all_docs: List[Document] = []
    processed_urls = set()

    for path in sorted(folder.rglob("*")):

        if not path.is_file():
            continue

        suffix = path.suffix.lower()

        if suffix not in supported_files:
            continue

        # ----------------------------------------------------
        # Windows .url / macOS .webloc
        # ----------------------------------------------------

        if suffix in {".url", ".webloc"}:

            url = extract_url_from_shortcut(path)

            if url and url not in processed_urls:

                processed_urls.add(url)

                print(f"[INFO] Fetching link: {url}")

                all_docs.extend(
                    fetch_url(url)
                )

            continue

        # ----------------------------------------------------
        # .links / .urls
        # ----------------------------------------------------

        if suffix in {".links", ".urls"}:

            urls = extract_urls_from_link_file(path)

            for url in urls:

                if url in processed_urls:
                    continue

                processed_urls.add(url)

                print(f"[INFO] Fetching link: {url}")

                all_docs.extend(
                    fetch_url(url)
                )

            continue

        # ----------------------------------------------------
        # Local PDF
        # ----------------------------------------------------

        if suffix == ".pdf":

            all_docs.extend(
                load_pdf_file(path)
            )

        # ----------------------------------------------------
        # Local DOCX
        # ----------------------------------------------------

        elif suffix == ".docx":

            all_docs.extend(
                load_docx_file(path)
            )

        # ----------------------------------------------------
        # Plain-text file
        # ----------------------------------------------------

        else:

            all_docs.extend(
                load_text_file(path)
            )

    all_docs = [
        doc
        for doc in all_docs
        if doc.page_content
        and doc.page_content.strip()
    ]

    if not all_docs:
        raise ValueError(
            f"No readable documents or links were found in:\n"
            f"{folder}\n\n"
            f"Supported formats:\n"
            f"PDF, DOCX, TXT, MD, CSV, JSON, HTML, "
            f".url, .webloc, .links, .urls"
        )

    print(
        f"[INFO] Loaded {len(all_docs)} source documents/pages."
    )

    return all_docs


# ============================================================
# INGESTION
# ============================================================

docs = load_documents_from_folder(
    DOCUMENTS_DIR
)


# ============================================================
# CHUNKING
# ============================================================

chunks = RecursiveCharacterTextSplitter(
    chunk_size=900,
    chunk_overlap=150,
).split_documents(docs)

for doc in chunks:
    doc.page_content = clean_text(doc.page_content)

chunks = [
    doc
    for doc in chunks
    if doc.page_content
]

if not chunks:
    raise ValueError(
        "No usable chunks were created from ./documents/"
    )

print(
    f"[INFO] Created {len(chunks)} chunks."
)


# ============================================================
# LOCAL BGE EMBEDDINGS
# ============================================================

print("[INFO] Loading BGE embeddings...")

embeddings = HuggingFaceEmbeddings(
    model_name="BAAI/bge-base-en-v1.5",
    model_kwargs={
        "device": "cpu",
    },
    encode_kwargs={
        "normalize_embeddings": True,
    },
)


# ============================================================
# FAISS
# ============================================================

print("[INFO] Building FAISS vector store...")

vector_store = FAISS.from_documents(
    chunks,
    embeddings,
)

retriever = vector_store.as_retriever(
    search_type="similarity",
    search_kwargs={
        "k": 4,
    },
)


# ============================================================
# GROQ GPT-OSS 20B
# ============================================================

print("[INFO] Initializing Groq GPT-OSS 20B...")

llm = ChatGroq(
    model="openai/gpt-oss-20b",
    temperature=0,
    max_tokens=4096,
    reasoning_effort="low",
    reasoning_format="hidden",
    max_retries=2,
)


# ============================================================
# CRAG THRESHOLDS
# ============================================================

UPPER_TH = 0.7
LOWER_TH = 0.3


# ============================================================
# STATE
# ============================================================

class State(TypedDict):
    question: str
    docs: List[Document]
    good_docs: List[Document]
    verdict: str
    reason: str
    strips: List[str]
    kept_strips: List[str]
    refined_context: str
    web_query: str
    web_docs: List[Document]
    answer: str


# ============================================================
# RETRIEVE NODE
# ============================================================

def retrieve_node(state: State) -> dict:

    question = state["question"]

    docs = retriever.invoke(
        question
    )

    return {
        "docs": docs
    }


# ============================================================
# DOCUMENT EVALUATION
# ============================================================

class DocEvalScore(BaseModel):
    score: float = Field(
        description="Relevance score from 0.0 to 1.0."
    )
    reason: str = Field(
        description="Short reason for the score."
    )


doc_eval_prompt = ChatPromptTemplate.from_messages(
    [
        (
            "system",
            "You are a strict retrieval evaluator for RAG.\n"
            "You will be given ONE retrieved chunk and a question.\n"
            "Score how relevant the chunk is to the question.\n"
            "0.0 = irrelevant.\n"
            "1.0 = sufficient to answer fully or mostly.\n"
            "Be conservative with high scores."
        ),
        (
            "human",
            "Question: {question}\n\n"
            "Chunk:\n{chunk}"
        ),
    ]
)


doc_eval_chain = (
    doc_eval_prompt
    | llm.with_structured_output(
        DocEvalScore,
        method="json_schema",
        strict=True,
    )
)


def eval_each_doc_node(state: State) -> dict:

    question = state["question"]

    scores: List[float] = []
    good_docs: List[Document] = []

    for doc in state["docs"]:

        try:

            result = doc_eval_chain.invoke(
                {
                    "question": question,
                    "chunk": doc.page_content,
                }
            )

            score = max(
                0.0,
                min(
                    1.0,
                    float(result.score),
                ),
            )

            scores.append(score)

            if score > LOWER_TH:
                good_docs.append(doc)

        except Exception as e:

            print(
                f"[WARNING] Document evaluation failed: {e}"
            )

            scores.append(0.0)

    if any(
        score > UPPER_TH
        for score in scores
    ):

        return {
            "good_docs": good_docs,
            "verdict": "CORRECT",
            "reason": (
                f"At least one retrieved chunk scored "
                f"> {UPPER_TH}."
            ),
        }

    if (
        scores
        and all(
            score < LOWER_TH
            for score in scores
        )
    ):

        return {
            "good_docs": [],
            "verdict": "INCORRECT",
            "reason": (
                f"All retrieved chunks scored "
                f"< {LOWER_TH}."
            ),
        }

    return {
        "good_docs": good_docs,
        "verdict": "AMBIGUOUS",
        "reason": (
            f"No chunk scored > {UPPER_TH}, "
            f"but not all were < {LOWER_TH}."
        ),
    }


# ============================================================
# SENTENCE DECOMPOSITION
# ============================================================

def decompose_to_sentences(
    text: str,
) -> List[str]:

    text = re.sub(
        r"\s+",
        " ",
        text,
    ).strip()

    sentences = re.split(
        r"(?<=[.!?])\s+",
        text,
    )

    return [
        sentence.strip()
        for sentence in sentences
        if len(sentence.strip()) > 20
    ]


# ============================================================
# SENTENCE FILTER
# ============================================================

class KeepOrDrop(BaseModel):
    keep: bool = Field(
        description=(
            "True only when the sentence directly "
            "helps answer the question."
        )
    )


filter_prompt = ChatPromptTemplate.from_messages(
    [
        (
            "system",
            "You are a strict relevance filter.\n"
            "Return keep=true only if the sentence directly "
            "helps answer the question.\n"
            "Do not use information outside the sentence."
        ),
        (
            "human",
            "Question: {question}\n\n"
            "Sentence:\n{sentence}"
        ),
    ]
)


filter_chain = (
    filter_prompt
    | llm.with_structured_output(
        KeepOrDrop,
        method="json_schema",
        strict=True,
    )
)


# ============================================================
# REFINE
# ============================================================

def refine(state: State) -> dict:

    question = state["question"]

    if state.get("verdict") == "CORRECT":

        docs_to_use = state["good_docs"]

    elif state.get("verdict") == "INCORRECT":

        docs_to_use = state["web_docs"]

    else:

        docs_to_use = (
            state["good_docs"]
            + state["web_docs"]
        )

    context = "\n\n".join(
        doc.page_content
        for doc in docs_to_use
        if doc.page_content
    ).strip()

    strips = decompose_to_sentences(
        context
    )

    kept: List[str] = []

    for sentence in strips:

        try:

            result = filter_chain.invoke(
                {
                    "question": question,
                    "sentence": sentence,
                }
            )

            if result.keep:
                kept.append(sentence)

        except Exception as e:

            print(
                f"[WARNING] Sentence filter failed: {e}"
            )

    refined_context = "\n".join(
        kept
    ).strip()

    return {
        "strips": strips,
        "kept_strips": kept,
        "refined_context": refined_context,
    }


# ============================================================
# WEB QUERY REWRITE
# ============================================================

class WebQuery(BaseModel):
    query: str = Field(
        description="Short web search query of 6 to 14 words."
    )


rewrite_prompt = ChatPromptTemplate.from_messages(
    [
        (
            "system",
            "Rewrite the user's question into a short web "
            "search query composed of useful keywords.\n"
            "Rules:\n"
            "- Keep it 6 to 14 words.\n"
            "- Preserve the main meaning.\n"
            "- If the question implies recency, preserve "
            "that recency constraint.\n"
            "- Do not answer the question."
        ),
        (
            "human",
            "Question: {question}"
        ),
    ]
)


rewrite_chain = (
    rewrite_prompt
    | llm.with_structured_output(
        WebQuery,
        method="json_schema",
        strict=True,
    )
)


def rewrite_query_node(state: State) -> dict:

    try:

        result = rewrite_chain.invoke(
            {
                "question": state["question"]
            }
        )

        return {
            "web_query": result.query.strip()
        }

    except Exception as e:

        print(
            f"[WARNING] Query rewrite failed: {e}"
        )

        return {
            "web_query": state["question"]
        }


# ============================================================
# TAVILY
# ============================================================

tavily = TavilySearch(
    max_results=5,
    topic="general",
    search_depth="basic",
)


# ============================================================
# WEB SEARCH
# ============================================================

def web_search_node(state: State) -> dict:

    query = (
        state.get("web_query")
        or state["question"]
    )

    try:

        raw_result = tavily.invoke(
            {
                "query": query
            }
        )

    except Exception as e:

        print(
            f"[WARNING] Tavily search failed: {e}"
        )

        return {
            "web_docs": []
        }

    # Current TavilySearch returns a dictionary
    # containing a "results" list.
    if isinstance(raw_result, dict):

        results = raw_result.get(
            "results",
            []
        )

    elif isinstance(raw_result, list):

        results = raw_result

    else:

        results = []

    web_docs: List[Document] = []

    for result in results:

        if not isinstance(result, dict):
            continue

        title = result.get(
            "title",
            ""
        )

        url = result.get(
            "url",
            ""
        )

        content = (
            result.get("content", "")
            or result.get("raw_content", "")
            or result.get("snippet", "")
        )

        if not content:
            continue

        page_content = (
            f"TITLE: {title}\n"
            f"URL: {url}\n"
            f"CONTENT:\n{content}"
        )

        web_docs.append(
            Document(
                page_content=page_content,
                metadata={
                    "source": url,
                    "url": url,
                    "title": title,
                    "type": "web_search",
                },
            )
        )

    print(
        f"[INFO] Tavily returned "
        f"{len(web_docs)} usable results."
    )

    return {
        "web_docs": web_docs
    }


# ============================================================
# ANSWER GENERATION
# ============================================================

answer_prompt = ChatPromptTemplate.from_messages(
    [
        (
            "system",
            "You are a helpful ML tutor.\n"
            "Answer ONLY using the provided context.\n"
            "Do not use outside knowledge.\n"
            "If the context is empty or insufficient, "
            "say exactly: I don't know."
        ),
        (
            "human",
            "Question: {question}\n\n"
            "Context:\n{context}"
        ),
    ]
)


def generate(state: State) -> dict:

    try:

        result = (
            answer_prompt
            | llm
        ).invoke(
            {
                "question": state["question"],
                "context": state["refined_context"],
            }
        )

        return {
            "answer": result.content
        }

    except Exception as e:

        return {
            "answer": (
                "Unable to generate the answer: "
                f"{e}"
            )
        }


# ============================================================
# ROUTING
# ============================================================

def route_after_eval(
    state: State,
) -> str:

    if state["verdict"] == "CORRECT":

        return "refine"

    return "rewrite_query"


# ============================================================
# LANGGRAPH
# ============================================================

graph = StateGraph(State)

graph.add_node(
    "retrieve",
    retrieve_node,
)

graph.add_node(
    "eval_each_doc",
    eval_each_doc_node,
)

graph.add_node(
    "rewrite_query",
    rewrite_query_node,
)

graph.add_node(
    "web_search",
    web_search_node,
)

graph.add_node(
    "refine",
    refine,
)

graph.add_node(
    "generate",
    generate,
)


graph.add_edge(
    START,
    "retrieve",
)

graph.add_edge(
    "retrieve",
    "eval_each_doc",
)

graph.add_conditional_edges(
    "eval_each_doc",
    route_after_eval,
    {
        "refine": "refine",
        "rewrite_query": "rewrite_query",
    },
)

graph.add_edge(
    "rewrite_query",
    "web_search",
)

graph.add_edge(
    "web_search",
    "refine",
)

graph.add_edge(
    "refine",
    "generate",
)

graph.add_edge(
    "generate",
    END,
)


# ============================================================
# COMPILE
# ============================================================

app = graph.compile()


# ============================================================
# INTERACTIVE RUN
# ============================================================

if __name__ == "__main__":

    print("\n" + "=" * 70)
    print("CRAG - Groq GPT-OSS 20B")
    print("Type your question.")
    print("Type 'exit' to stop.")
    print("=" * 70)

    while True:

        question = input("\nQuestion: ").strip()

        if question.lower() in {
            "exit",
            "quit",
            "q",
        }:
            print("\nExiting...")
            break

        if not question:
            print("Please enter a question.")
            continue

        try:

            result = app.invoke(
                {
                    "question": question,
                    "docs": [],
                    "good_docs": [],
                    "verdict": "",
                    "reason": "",
                    "strips": [],
                    "kept_strips": [],
                    "refined_context": "",
                    "web_query": "",
                    "web_docs": [],
                    "answer": "",
                }
            )

            print("\n" + "-" * 70)

            print(
                "VERDICT:",
                result["verdict"],
            )

            print(
                "REASON:",
                result["reason"],
            )

            if result.get("web_query"):
                print(
                    "WEB_QUERY:",
                    result["web_query"],
                )

            print("\nANSWER:\n")
            print(
                result["answer"]
            )

            print("-" * 70)

        except Exception as e:

            print(
                f"\n[ERROR] {e}"
            )
